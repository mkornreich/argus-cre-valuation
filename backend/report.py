"""ARGUS-style report pack: an .xlsx workbook built from an analysis result.

Sheets: Summary, Cash Flow (annual), Cash Flow (monthly), Rent Roll,
Sensitivity, Traditional Valuation — the core of an ARGUS report set.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import engine

_NAVY = "0F2A43"
_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor=_NAVY)
_BOLD = Font(bold=True)
_MONEY = "#,##0"
_PCT = "0.00%"
_THIN = Border(bottom=Side(style="thin", color="DDDDDD"))

# Proforma line items: (label, result key, is_bold)
_LINES = [
    ("Potential Base Rent", "potential_base_rent", False),
    ("Absorption & Turnover Vacancy", "absorption_turnover_vacancy", False),
    ("Base Rent Abatements", "base_rent_abatements", False),
    ("Scheduled Base Rent", "scheduled_base_rent", True),
    ("Percentage Rent", "percentage_rent", False),
    ("Expense Reimbursements", "expense_reimbursements", False),
    ("Total Tenant Revenue", "total_tenant_revenue", True),
    ("General Vacancy & Credit Loss", "general_vacancy", False),
    ("Effective Gross Revenue", "effective_gross_revenue", True),
    ("Operating Expenses", "operating_expenses", False),
    ("Net Operating Income", "noi", True),
    ("Tenant Improvements", "tenant_improvements", False),
    ("Leasing Commissions", "leasing_commissions", False),
    ("Capital Reserve", "capital_reserve", False),
    ("Unlevered Cash Flow", "unlevered_cash_flow", True),
    ("Debt Service", "debt_service", False),
    ("Levered Cash Flow", "levered_cash_flow", True),
]


def _title(ws, text):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=_NAVY)


def _header_row(ws, row, values):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = _HEAD
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="right" if c > 1 else "left")


def _proforma_sheet(ws, title, rows):
    _title(ws, title)
    periods = [r.get("period", r.get("year")) for r in rows]
    _header_row(ws, 3, ["Line item", *periods])
    for i, (label, key, bold) in enumerate(_LINES, start=4):
        if not any(key in r for r in rows):
            continue
        ws.cell(row=i, column=1, value=label).font = _BOLD if bold else Font()
        for c, r in enumerate(rows, start=2):
            cell = ws.cell(row=i, column=c, value=r.get(key))
            cell.number_format = _MONEY
            if bold:
                cell.font = _BOLD
    ws.column_dimensions["A"].width = 32
    for c in range(2, len(periods) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 13


def build_report(scenario: dict) -> io.BytesIO:
    r = engine.analyze(scenario, with_sensitivity=True, include_monthly=True)
    m = r["metrics"]
    tv = r["traditional_valuation"]
    wb = Workbook()

    # --- Summary -------------------------------------------------------- #
    ws = wb.active
    ws.title = "Summary"
    _title(ws, r["property"]["name"])
    ws["A2"] = f"{r['property']['property_type']} · {r['property']['total_sf']:,.0f} SF"
    rows = [
        ("Purchase Price", m["purchase_price"], _MONEY),
        ("Price / SF", m["price_per_sf"], _MONEY),
        ("Year-1 NOI", m["year1_noi"], _MONEY),
        ("Going-in Cap Rate", m["going_in_cap_rate"], _PCT),
        ("Discount Rate", m["discount_rate"], _PCT),
        ("Exit Cap Rate", m["exit_cap_rate"], _PCT),
        ("Terminal Value", m["terminal_value"], _MONEY),
        ("DCF Value", m["dcf_value"], _MONEY),
        ("NPV vs. Price", m["npv"], _MONEY),
        ("Unlevered IRR", m["unlevered_irr"], _PCT),
        ("Unlevered Equity Multiple", m["equity_multiple_unlevered"], "0.00"),
        ("Total Leasing & Capital", m["total_leasing_capital"], _MONEY),
    ]
    if m.get("debt"):
        d = m["debt"]
        rows += [("Loan Amount", d["loan_amount"], _MONEY), ("Equity Invested", d["equity_invested"], _MONEY),
                 ("LTV", d["ltv_pct"], _PCT), ("DSCR (Yr 1)", d["dscr_year1"], "0.00"),
                 ("Levered IRR", d["levered_irr"], _PCT),
                 ("Levered Equity Multiple", d["equity_multiple_levered"], "0.00"),
                 ("Avg Cash-on-Cash", d["avg_cash_on_cash"], _PCT)]
    for i, (label, val, fmt) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = _BOLD
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18

    _proforma_sheet(wb.create_sheet("Cash Flow (Annual)"), "Operating Proforma — Annual", r["cash_flow"])
    _proforma_sheet(wb.create_sheet("Cash Flow (Monthly)"), "Operating Proforma — Monthly",
                    r.get("monthly_cash_flow", []))

    # --- Rent Roll ------------------------------------------------------ #
    rr = wb.create_sheet("Rent Roll")
    _title(rr, "Rent Roll")
    _header_row(rr, 3, ["Tenant", "SF", "Start Yr", "Term", "Rent/SF", "Review",
                        "Reimbursement", "Free (mo)"])
    for i, l in enumerate(engine.normalize_scenario(scenario)["leases"], start=4):
        vals = [l["tenant"], l["suite_sf"], l["start_year"], l["term_years"], l["base_rent_psf"],
                l["escalation_type"], l["reimbursement_method"], l["free_rent_months"]]
        for c, v in enumerate(vals, start=1):
            rr.cell(row=i, column=c, value=v)
    rr.column_dimensions["A"].width = 30

    # --- Sensitivity ---------------------------------------------------- #
    sv = wb.create_sheet("Sensitivity")
    _title(sv, "DCF Value — Discount Rate (rows) × Exit Cap (cols)")
    s = r["sensitivity"]
    _header_row(sv, 3, ["Discount \\ Exit", *[f"{c:.2%}" for c in s["exit_cap_rates"]]])
    for i, d in enumerate(s["discount_rates"], start=4):
        sv.cell(row=i, column=1, value=f"{d:.2%}").font = _BOLD
        for c, v in enumerate(s["dcf_value"][i - 4], start=2):
            cell = sv.cell(row=i, column=c, value=v)
            cell.number_format = _MONEY
    sv.column_dimensions["A"].width = 16

    # --- Traditional Valuation ----------------------------------------- #
    tvs = wb.create_sheet("Traditional Valuation")
    _title(tvs, "Traditional Valuation")
    trows = [
        ("Net Initial Yield", tv["net_initial_yield"], _PCT),
        ("Reversionary Yield", tv["reversionary_yield"], _PCT),
        ("Equivalent Yield", tv["equivalent_yield"], _PCT),
        ("Passing NOI", tv["passing_noi"], _MONEY),
        ("Reversion (ERV) NOI", tv["reversion_noi"], _MONEY),
        ("Years to Reversion (WALT)", tv["years_to_reversion"], "0.00"),
        ("Term Value", tv["term_value"], _MONEY),
        ("Reversion Value", tv["reversion_value"], _MONEY),
        ("Term & Reversion Value", tv["term_and_reversion_value"], _MONEY),
        ("Hardcore / Layer Value", tv["hardcore_value"], _MONEY),
    ]
    for i, (label, val, fmt) in enumerate(trows, start=4):
        tvs.cell(row=i, column=1, value=label).font = _BOLD
        tvs.cell(row=i, column=2, value=val).number_format = fmt
    tvs.column_dimensions["A"].width = 28
    tvs.column_dimensions["B"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
